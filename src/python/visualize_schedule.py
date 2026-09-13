"""Export a generated schedule JSON file as a readable Excel workbook."""

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = REPO / "src" / "renderer" / "src" / "components" / "schedule.json"
DEFAULT_OUTPUT = REPO / "schedule_table.xlsx"

SERVICE_LABELS = {
    "Stroke": "STK",
    "VA": "VA",
    "UH": "UH",
    "ELECTIVE": "EL",
    "CC": "CC",
    "VAC": "VAC",
    "NF": "NF",
    "EEG": "EEG",
    "B/U": "B/U",
    "NICU": "NICU",
    "CHILD": "CHLD",
    "CLINIC": "CLIN",
    "RAD": "RAD",
    "NFCL": "NFCL",
    "CONSULTS": "CONS",
    "EMG": "EMG",
    "EMU": "EMU",
    "JEOPARDY-ELECTIVE": "JEOP",
    "UNSCHEDULED": "UNSC",
    "": "",
}

COLOR_MAP = {
    "Stroke": "90EE90",
    "VA": "800080",
    "UH": "FFFF00",
    "ELECTIVE": "D3D3D3",
    "CC": "000000",
    "VAC": "FF0000",
    "NF": "ADD8E6",
    "EEG": "FFD700",
    "B/U": "8B4513",
    "NICU": "FFA07A",
    "CHILD": "FFB6C1",
    "CLINIC": "20B2AA",
    "RAD": "87CEFA",
    "NFCL": "9370DB",
    "CONSULTS": "FFE4B5",
    "EMG": "7FFFD4",
    "EMU": "FFEFD5",
    "JEOPARDY-ELECTIVE": "98FB98",
    "UNSCHEDULED": "FFA500",
    "": "FFFFFF",
}

WHITE_TEXT = {"CC", "VAC", "VA", "B/U", "NFCL"}


def load_schedule(path: Path):
    with path.open("r", encoding="utf-8") as schedule_file:
        data = json.load(schedule_file)

    weekly_counts = data.pop("weekly_counts", None)
    rows = [(name, list(weeks)) for name, weeks in data.items()]
    num_weeks = len(weekly_counts) if weekly_counts is not None else max(
        (len(weeks) for _, weeks in rows), default=0
    )

    normalized_rows = []
    for name, weeks in rows:
        normalized_rows.append((name, weeks[:num_weeks] + [None] * (num_weeks - len(weeks))))

    if weekly_counts is not None:
        weekly_counts = list(weekly_counts[:num_weeks])
        weekly_counts.extend([None] * (num_weeks - len(weekly_counts)))

    return normalized_rows, weekly_counts


def autosize(worksheet):
    for column_cells in worksheet.columns:
        max_length = max(
            len("" if cell.value is None else str(cell.value))
            for cell in column_cells
        )
        column_letter = get_column_letter(column_cells[0].column)
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 30)


def write_workbook(rows, weekly_counts, output_path: Path):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    num_weeks = len(rows[0][1]) if rows else len(weekly_counts or [])

    workbook = Workbook()
    schedule_sheet = workbook.active
    schedule_sheet.title = "Schedule"
    schedule_sheet.append(["Resident", *[f"Week {week}" for week in range(1, num_weeks + 1)]])

    for resident, assignments in rows:
        schedule_sheet.append([resident, *assignments])

    weekly_count_row = None
    if weekly_counts is not None:
        schedule_sheet.append(["weekly_counts", *weekly_counts])
        weekly_count_row = schedule_sheet.max_row

    schedule_sheet.freeze_panes = "B2"
    for row in schedule_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in schedule_sheet.iter_rows(min_row=2, min_col=2):
        for cell in row:
            if weekly_count_row is not None and cell.row == weekly_count_row:
                continue
            service_name = cell.value if isinstance(cell.value, str) else ""
            cell.value = SERVICE_LABELS.get(service_name, service_name)
            fill_hex = COLOR_MAP.get(service_name)
            if fill_hex:
                cell.fill = PatternFill(fill_type="solid", fgColor=fill_hex)
            if service_name in WHITE_TEXT:
                cell.font = Font(color="FFFFFF")

    if weekly_count_row is not None:
        for cell in schedule_sheet[weekly_count_row]:
            cell.fill = PatternFill(fill_type="solid", fgColor="D3D3D3")
            cell.font = Font(bold=True)

    autosize(schedule_sheet)

    legend_sheet = workbook.create_sheet("Legend")
    legend_sheet.append(["Service", "Label", "Color"])
    for service_name, fill_hex in COLOR_MAP.items():
        legend_sheet.append([service_name, SERVICE_LABELS.get(service_name, service_name), f"#{fill_hex}"])
        legend_sheet.cell(row=legend_sheet.max_row, column=3).fill = PatternFill(
            fill_type="solid", fgColor=fill_hex
        )
    autosize(legend_sheet)

    workbook.save(output_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Schedule JSON path")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Excel workbook path")
    args = parser.parse_args()

    rows, weekly_counts = load_schedule(args.input)
    write_workbook(rows, weekly_counts, args.output)
    print(f"Colored schedule saved to '{args.output}'")


if __name__ == "__main__":
    main()
