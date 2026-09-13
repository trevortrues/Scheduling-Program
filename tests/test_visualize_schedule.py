import importlib.util
import tempfile
import unittest
from pathlib import Path


REPO = Path(__file__).resolve().parents[1]
VISUALIZER_PATH = REPO / "src" / "python" / "visualize_schedule.py"
SCHEDULE_PATH = REPO / "src" / "renderer" / "src" / "components" / "schedule.json"
OPENPYXL_AVAILABLE = importlib.util.find_spec("openpyxl") is not None


@unittest.skipUnless(OPENPYXL_AVAILABLE, "Install requirements.txt to run visualizer tests")
class VisualizerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location("schedule_visualizer", VISUALIZER_PATH)
        cls.visualizer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.visualizer)

    def test_exports_repository_schedule_to_workbook(self):
        rows, weekly_counts = self.visualizer.load_schedule(SCHEDULE_PATH)

        self.assertEqual(len(rows), 27)
        self.assertEqual(len(rows[0][1]), 53)
        self.assertEqual(len(weekly_counts), 53)

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "schedule_table.xlsx"
            self.visualizer.write_workbook(rows, weekly_counts, output_path)

            from openpyxl import load_workbook

            workbook = load_workbook(output_path)
            self.assertEqual(workbook.sheetnames, ["Schedule", "Legend"])
            self.assertEqual(workbook["Schedule"].freeze_panes, "B2")
            self.assertEqual(workbook["Schedule"].max_column, 54)
            self.assertEqual(workbook["Legend"]["A1"].value, "Service")


if __name__ == "__main__":
    unittest.main()
